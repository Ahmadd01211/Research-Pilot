"""Bounded export generation stores exact download bytes for history."""

import csv
import hashlib
import io
import json
from datetime import UTC, datetime
from uuid import UUID

from docx import Document
from fastapi import HTTPException
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.modules.exports.models import ExportRecord
from app.modules.exports.schemas import ExportInput, ExportPage, ExportRead
from app.modules.projects.service import owned as owned_project
from app.modules.references.models import ReferenceRecord
from app.modules.writing.models import WritingDraft

MAX_EXPORT_BYTES = 100 * 1024 * 1024
MAX_FILE_BYTES = 2 * 1024 * 1024


def _read(item: ExportRecord) -> ExportRead:
    return ExportRead(id=item.id, kind=item.kind, draft_id=item.draft_id,
                      filename=item.filename, sha256=item.sha256,
                      byte_size=item.byte_size, created_at=item.created_at)


def _safe_cell(value: object) -> str:
    text = str(value) if value is not None else ""
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


REFERENCE_HEADERS = ["Type", "Title", "Authors", "Year", "Journal or container", "DOI", "URL", "Linked PDF ID"]


def _references(db: Session, project_id: UUID) -> list[ReferenceRecord]:
    return list(db.scalars(select(ReferenceRecord).where(
        ReferenceRecord.project_id == project_id, ReferenceRecord.deleted_at.is_(None),
    ).order_by(ReferenceRecord.title_key, ReferenceRecord.id)).all())


def _reference_values(record: ReferenceRecord) -> tuple[object, ...]:
    return (record.kind, record.title, "; ".join(record.authors), record.year,
            record.container_title, record.doi, record.url, record.document_id)


def _reference_csv(db: Session, project_id: UUID) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(REFERENCE_HEADERS)
    for record in _references(db, project_id):
        writer.writerow([_safe_cell(item) for item in _reference_values(record)])
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _reference_xlsx(db: Session, project_id: UUID) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "References"
    sheet.append(REFERENCE_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for row, record in enumerate(_references(db, project_id), start=2):
        for column, value in enumerate(_reference_values(record), 1):
            cell = sheet.cell(row, column)
            if column == 4 and value is not None:
                cell.value = value
            else:
                cell.value = str(value) if value is not None else ""
                cell.data_type = "s"  # Preserve literal metadata; never emit formulas.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for letter, width in zip("ABCDEFGH", (18, 52, 48, 12, 44, 32, 52, 40), strict=True):
        sheet.column_dimensions[letter].width = width
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _draft(db: Session, project_id: UUID, data: ExportInput) -> WritingDraft:
    if data.draft_id is None or data.expected_revision is None:
        raise HTTPException(422, "Choose a draft and its current revision")
    draft = db.scalar(select(WritingDraft).where(
        WritingDraft.id == data.draft_id, WritingDraft.project_id == project_id,
        WritingDraft.deleted_at.is_(None),
    ))
    if draft is None:
        raise HTTPException(404, "Draft not found")
    if draft.revision != data.expected_revision:
        raise HTTPException(409, "Draft changed. Review its latest version before exporting")
    return draft


def _draft_markdown(draft: WritingDraft) -> bytes:
    title = " ".join(draft.title.splitlines())
    return f"# {title}\n\n{draft.body.rstrip()}\n".encode()


def _draft_docx(draft: WritingDraft) -> bytes:
    document = Document()
    title = " ".join(draft.title.splitlines())
    document.core_properties.title = title
    document.add_heading(title, level=0)
    for line in draft.body.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        else:
            document.add_paragraph(line)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def create(db: Session, owner: UUID, project_id: UUID, data: ExportInput) -> ExportRead:
    project = owned_project(db, owner, project_id, lock=True)
    fingerprint = hashlib.sha256(json.dumps(data.model_dump(mode="json"), sort_keys=True).encode()).hexdigest()
    existing = db.scalar(select(ExportRecord).where(
        ExportRecord.project_id == project_id, ExportRecord.request_id == data.request_id,
    ))
    if existing:
        if existing.request_hash != fingerprint:
            raise HTTPException(409, "This export request ID was already used")
        return _read(existing)
    if project.archived_at is not None:
        raise HTTPException(409, "Restore the project before creating exports")
    if data.kind in {"reference_csv", "reference_xlsx"}:
        if data.draft_id is not None or data.expected_revision is not None:
            raise HTTPException(422, "Reference exports do not use a draft")
        extension = "csv" if data.kind == "reference_csv" else "xlsx"
        content = (_reference_csv(db, project_id) if extension == "csv"
                   else _reference_xlsx(db, project_id))
        filename = f"references-{project_id.hex[:8]}.{extension}"
        mime = ("text/csv; charset=utf-8" if extension == "csv" else
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        draft = _draft(db, project_id, data)
        extension = "md" if data.kind == "draft_markdown" else "docx"
        content = _draft_markdown(draft) if extension == "md" else _draft_docx(draft)
        assert data.draft_id is not None
        filename = f"draft-{data.draft_id.hex[:8]}-v{data.expected_revision}.{extension}"
        mime = ("text/markdown; charset=utf-8" if extension == "md" else
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    if not content or len(content) > MAX_FILE_BYTES:
        raise HTTPException(413, "Export exceeds the 2 MB file limit")
    used = db.scalar(select(func.sum(ExportRecord.byte_size)).where(
        ExportRecord.project_id == project_id,
    )) or 0
    if used + len(content) > MAX_EXPORT_BYTES:
        raise HTTPException(409, "Project export history reached its 100 MB storage limit")
    item = ExportRecord(project_id=project_id, draft_id=data.draft_id, kind=data.kind,
                        filename=filename, content_type=mime, body=content,
                        sha256=hashlib.sha256(content).hexdigest(), byte_size=len(content),
                        request_id=data.request_id, request_hash=fingerprint,
                        created_at=datetime.now(UTC))
    db.add(item)
    db.commit()
    return _read(item)


def list_exports(db: Session, owner: UUID, project_id: UUID, offset: int) -> ExportPage:
    owned_project(db, owner, project_id)
    total = db.scalar(select(func.count()).select_from(ExportRecord).where(
        ExportRecord.project_id == project_id,
    )) or 0
    items = db.scalars(select(ExportRecord).where(
        ExportRecord.project_id == project_id,
    ).order_by(ExportRecord.created_at.desc(), ExportRecord.id.desc()).offset(offset).limit(20)).all()
    return ExportPage(items=[_read(item) for item in items], total=total)


def download(db: Session, owner: UUID, project_id: UUID, export_id: UUID) -> ExportRecord:
    owned_project(db, owner, project_id)
    item = db.scalar(select(ExportRecord).where(
        ExportRecord.id == export_id, ExportRecord.project_id == project_id,
    ))
    if item is None:
        raise HTTPException(404, "Export not found")
    return item
