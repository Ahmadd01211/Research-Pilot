"""Exports are immutable downloads with owner checks and spreadsheet-safe CSV."""

import io
from uuid import uuid4

from docx import Document
from openpyxl import load_workbook

from tests.conftest import csrf, login, register
from tests.test_matrix import setup_source


def test_reference_csv_snapshot_is_safe_and_immutable(client):
    project, document_id = setup_source(client)
    base = f"/api/v1/projects/{project['id']}"
    reference = client.post(base + "/references", headers=csrf(client), json={
        "kind": "journal", "title": "=SUM(1,1)", "authors": ["A. Scholar"],
        "year": 2024, "container_title": "Research Journal", "doi": None,
        "url": "", "document_id": document_id,
    })
    assert reference.status_code == 201, reference.text
    request = {"kind": "reference_csv", "request_id": str(uuid4())}
    created = client.post(base + "/exports", headers=csrf(client), json=request)
    assert created.status_code == 201, created.text
    item = created.json()
    download = client.get(base + "/exports/" + item["id"] + "/download")
    assert download.status_code == 200
    assert b"'=SUM(1,1)" in download.content
    assert download.headers["content-disposition"].startswith("attachment;")
    assert client.post(base + "/exports", headers=csrf(client), json=request).json()["id"] == item["id"]
    assert client.post(base + "/exports", headers=csrf(client), json={
        **request, "kind": "draft_markdown",
    }).status_code == 409
    ref = reference.json()
    assert client.put(base + "/references/" + ref["id"], headers=csrf(client), json={
        "kind": "journal", "title": "Safe revised title", "authors": ["A. Scholar"],
        "year": 2024, "container_title": "Research Journal", "doi": None,
        "url": "", "document_id": document_id, "expected_revision": 1,
    }).status_code == 200
    assert client.get(base + "/exports/" + item["id"] + "/download").content == download.content
    assert client.get(base + "/exports").json()["items"][0]["id"] == item["id"]


def test_draft_export_rejects_stale_revision_and_foreign_access(client):
    project, _ = setup_source(client)
    base = f"/api/v1/projects/{project['id']}"
    draft = client.post(base + "/drafts", headers=csrf(client), json={
        "title": "Review", "body": "# Introduction\nResearcher text.",
    }).json()
    export = {"kind": "draft_markdown", "request_id": str(uuid4()),
              "draft_id": draft["id"], "expected_revision": 1}
    created = client.post(base + "/exports", headers=csrf(client), json=export)
    assert created.status_code == 201, created.text
    item = created.json()
    assert b"Researcher text." in client.get(base + "/exports/" + item["id"] + "/download").content
    assert client.put(base + "/drafts/" + draft["id"], headers=csrf(client), json={
        "title": "Revised", "body": "Revised text", "expected_revision": 1,
    }).status_code == 200
    assert client.post(base + "/exports", headers=csrf(client), json={
        **export, "request_id": str(uuid4()),
    }).status_code == 409
    register(client, "export-other@example.com")
    login(client, "export-other@example.com")
    assert client.get(base + "/exports").status_code == 404
    assert client.get(base + "/exports/" + item["id"] + "/download").status_code == 404


def test_xlsx_reference_snapshot_preserves_literal_cells_and_original_bytes(client):
    project, _ = setup_source(client)
    base = f"/api/v1/projects/{project['id']}"
    reference = client.post(base + "/references", headers=csrf(client), json={
        "kind": "journal", "title": "=SUM(1,1)", "authors": ["Scholar, Ada"],
        "year": 2024, "container_title": "Research Journal", "doi": None,
        "url": "", "document_id": None,
    })
    assert reference.status_code == 201, reference.text
    payload = {"kind": "reference_xlsx", "request_id": str(uuid4())}
    created = client.post(base + "/exports", headers=csrf(client), json=payload)
    assert created.status_code == 201, created.text
    item = created.json()
    assert item["filename"].endswith(".xlsx")
    downloaded = client.get(base + "/exports/" + item["id"] + "/download")
    assert downloaded.status_code == 200
    sheet = load_workbook(io.BytesIO(downloaded.content), read_only=False).active
    assert sheet["B2"].value == "=SUM(1,1)"
    assert sheet["B2"].data_type == "s"
    assert sheet["D2"].value == 2024
    assert sheet.freeze_panes == "A2"
    assert client.post(base + "/exports", headers=csrf(client), json=payload).json()["id"] == item["id"]
    assert client.put(base + "/references/" + reference.json()["id"], headers=csrf(client), json={
        "kind": "journal", "title": "Updated", "authors": ["Scholar, Ada"],
        "year": 2024, "container_title": "Research Journal", "doi": None,
        "url": "", "document_id": None, "expected_revision": 1,
    }).status_code == 200
    assert client.get(base + "/exports/" + item["id"] + "/download").content == downloaded.content


def test_docx_draft_snapshot_is_readable_and_revision_checked(client):
    project, _ = setup_source(client)
    base = f"/api/v1/projects/{project['id']}"
    draft = client.post(base + "/drafts", headers=csrf(client), json={
        "title": "Review", "body": "# Introduction\nResearcher text.",
    }).json()
    payload = {"kind": "draft_docx", "request_id": str(uuid4()),
               "draft_id": draft["id"], "expected_revision": 1}
    created = client.post(base + "/exports", headers=csrf(client), json=payload)
    assert created.status_code == 201, created.text
    item = created.json()
    assert item["filename"].endswith(".docx")
    downloaded = client.get(base + "/exports/" + item["id"] + "/download")
    assert downloaded.status_code == 200
    document = Document(io.BytesIO(downloaded.content))
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Review", "Introduction", "Researcher text.",
    ]
    assert document.paragraphs[1].style.name == "Heading 1"
    assert client.put(base + "/drafts/" + draft["id"], headers=csrf(client), json={
        "title": "Updated", "body": "New body", "expected_revision": 1,
    }).status_code == 200
    assert client.post(base + "/exports", headers=csrf(client), json={
        **payload, "request_id": str(uuid4()),
    }).status_code == 409
    assert client.get(base + "/exports/" + item["id"] + "/download").content == downloaded.content

def test_matrix_snapshots_preserve_notes_provenance_and_history(client):
    project, document_id = setup_source(client)
    base = f"/api/v1/projects/{project['id']}"
    row = client.post(base + "/matrix", headers=csrf(client), json={
        "document_id": document_id,
        "cells": [{"field": "method", "value": "=Researcher note", "chunk_id": "matrix-source-1"}],
    })
    assert row.status_code == 201, row.text
    downloads = []
    for kind in ("matrix_csv", "matrix_xlsx"):
        payload = {"kind": kind, "request_id": str(uuid4())}
        created = client.post(base + "/exports", headers=csrf(client), json=payload)
        assert created.status_code == 201, created.text
        item = created.json()
        response = client.get(base + "/exports/" + item["id"] + "/download")
        downloads.append((item["id"], response.content))
        assert client.post(base + "/exports", headers=csrf(client), json=payload).json()["id"] == item["id"]
        if kind == "matrix_csv":
            assert b"'=Researcher note" in response.content
            assert b"matrix-source-1" in response.content
            assert b"passage-1" in response.content
        else:
            workbook = load_workbook(io.BytesIO(response.content))
            assert workbook["Literature matrix"]["D2"].value == "=Researcher note"
            assert workbook["Literature matrix"]["D2"].data_type == "s"
            assert workbook["Source provenance"]["G2"].value == 2
            assert workbook["Source provenance"]["H2"].value == "matrix-source-1"
            assert workbook["Source provenance"]["I2"].value == row.json()["cells"][0]["source_sha256"]
    assert client.post(base + "/matrix/" + row.json()["id"] + "/remove",
                       headers=csrf(client), json={"expected_revision": 1}).status_code == 204
    for export_id, content in downloads:
        assert client.get(base + "/exports/" + export_id + "/download").content == content
    register(client, "matrix-export-foreign@example.com")
    login(client, "matrix-export-foreign@example.com")
    assert client.post(base + "/exports", headers=csrf(client), json={
        "kind": "matrix_xlsx", "request_id": str(uuid4()),
    }).status_code == 404

def test_matrix_xlsx_rejects_silent_metadata_truncation_but_csv_keeps_it(client):
    from app.modules.documents.models import EvidenceChunk

    project, document_id = setup_source(client)
    base = f"/api/v1/projects/{project['id']}"
    with client.app.state.session_factory() as db:
        chunk = db.get(EvidenceChunk, "matrix-source-1")
        chunk.source_metadata = {"annotation": "a" * 33000}
        db.commit()
    assert client.post(base + "/matrix", headers=csrf(client), json={
        "document_id": document_id,
        "cells": [{"field": "method", "value": "Survey", "chunk_id": "matrix-source-1"}],
    }).status_code == 201
    response = client.post(base + "/exports", headers=csrf(client), json={
        "kind": "matrix_xlsx", "request_id": str(uuid4()),
    })
    assert response.status_code == 413
    csv_export = client.post(base + "/exports", headers=csrf(client), json={
        "kind": "matrix_csv", "request_id": str(uuid4()),
    })
    assert csv_export.status_code == 201
    download = client.get(base + "/exports/" + csv_export.json()["id"] + "/download")
    assert b"a" * 33000 in download.content
